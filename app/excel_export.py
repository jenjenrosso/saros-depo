import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.database import db_session

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "kaynak", "ambar_yerlesim.xlsx")

def export_db_to_excel(output_path: str):
    """
    Veritabanindaki tum guncel stok ve malzemeleri orijinal sablon formatinda Excel'e aktarir.
    """
    if os.path.exists(TEMPLATE_PATH):
        wb = openpyxl.load_workbook(TEMPLATE_PATH)
    else:
        wb = openpyxl.Workbook()

    with db_session() as conn:
        lokasyonlar = conn.execute("SELECT * FROM lokasyonlar ORDER BY id").fetchall()

        for lok in lokasyonlar:
            sheet_name = lok["excel_sayfa"]
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.create_sheet(title=sheet_name)

            # Baslik satirini bul
            header_row_idx = None
            col_map = {}
            for r in range(1, 15):
                row_vals = [str(c.value).strip() if c.value else "" for c in ws[r]]
                if any("MALZEMENİN ADI" in v for v in row_vals):
                    header_row_idx = r
                    for c_idx, val in enumerate(row_vals):
                        if val:
                            col_map[val.strip()] = c_idx + 1 # 1-based openpyxl
                    break

            if not header_row_idx:
                continue

            # DB'den bu lokasyondaki malzemeleri cek
            malzemeler = conn.execute("""
                SELECT * FROM malzemeler WHERE lokasyon_id = ? ORDER BY id ASC
            """, (lok["id"],)).fetchall()

            # Malzemeleri yaz
            curr_row = header_row_idx + 1
            for idx, m in enumerate(malzemeler, 1):
                # Eger mevcut satir varsa uzerine yaz, yoksa yeni hucrelere yaz
                if col_map.get("S.NO"):
                    ws.cell(row=curr_row, column=col_map["S.NO"], value=idx)
                elif col_map.get("NO"):
                    ws.cell(row=curr_row, column=col_map["NO"], value=idx)

                if col_map.get("MALZEMENİN ADI"):
                    ws.cell(row=curr_row, column=col_map["MALZEMENİN ADI"], value=m["ad"])
                if col_map.get("MARKA"):
                    ws.cell(row=curr_row, column=col_map["MARKA"], value=m["marka"])
                if col_map.get("TİPİ-AÇIKLAMA"):
                    ws.cell(row=curr_row, column=col_map["TİPİ-AÇIKLAMA"], value=m["tip_aciklama"])

                if col_map.get("ADET"):
                    qty_val = m["miktar"]
                    if m["birim"] and m["birim"] != "adet":
                        qty_val = f"{qty_val} {m['birim']}"
                    ws.cell(row=curr_row, column=col_map["ADET"], value=qty_val)

                if col_map.get("RAF NO"):
                    ws.cell(row=curr_row, column=col_map["RAF NO"], value=m["raf_no"])
                if col_map.get("GİRİŞ T."):
                    ws.cell(row=curr_row, column=col_map["GİRİŞ T."], value=m["giris_tarihi"])
                if col_map.get("ÇIKIŞ T."):
                    ws.cell(row=curr_row, column=col_map["ÇIKIŞ T."], value=m["cikis_tarihi"])
                if col_map.get("RAF ÖMRÜ"):
                    ws.cell(row=curr_row, column=col_map["RAF ÖMRÜ"], value=m["raf_omru"])

                curr_row += 1

    wb.save(output_path)
    return output_path
